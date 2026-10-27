"""Exportación de resultados a CSV, Excel y gráficas PNG."""

import logging
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.styles import PatternFill

from src.engine import ResultadoSimulacion
from src.metrics import calcular_metricas
from src.scenarios import ResultadosPorCombinacion

logger = logging.getLogger(__name__)


def _resumen(resultados: ResultadosPorCombinacion) -> pd.DataFrame:
    filas: list[dict[str, Any]] = []
    metricas_promediables = (
        "makespan",
        "desviacion_estandar_carga",
        "tiempo_espera_promedio",
        "tareas_reasignadas",
        "sobrecargas_total",
        "tiempo_espera_cola_promedio",
        "tiempo_espera_cola_maximo",
    )
    for (escenario, estrategia), valor in resultados.items():
        ejecuciones = valor if isinstance(valor, list) else [valor]
        datos = [calcular_metricas(resultado) for resultado in ejecuciones]
        resumen = {
            "escenario": escenario,
            "estrategia": estrategia,
            "ejecuciones": len(ejecuciones),
            "tareas_completadas": min(dato["tareas_completadas"] for dato in datos),
        }
        for metrica in metricas_promediables:
            valores = [dato[metrica] for dato in datos]
            resumen[f"{metrica}_promedio"] = sum(valores) / len(valores)
            resumen[f"{metrica}_desviacion"] = (
                float(pd.Series(valores).std(ddof=0)) if len(valores) > 1 else 0.0
            )
        max_servidor = max(
            (len(dato["carga_por_servidor"]) for dato in datos), default=0
        )
        for servidor_id in range(max_servidor):
            resumen[f"tareas_servidor_{servidor_id}_promedio"] = sum(
                dato["tareas_por_servidor"].get(servidor_id, 0) for dato in datos
            ) / len(datos)
            resumen[f"carga_servidor_{servidor_id}_promedio"] = sum(
                dato["carga_por_servidor"].get(servidor_id, 0) for dato in datos
            ) / len(datos)
            resumen[f"sobrecargas_servidor_{servidor_id}_promedio"] = sum(
                dato["sobrecargas_por_servidor"].get(servidor_id, 0)
                for dato in datos
            ) / len(datos)
        resumen["sobrecargas_por_servidor_promedio"] = "; ".join(
            f"S{servidor_id}: "
            f"{resumen[f'sobrecargas_servidor_{servidor_id}_promedio']:.2f}"
            for servidor_id in range(max_servidor)
        )
        filas.append(resumen)
    tabla = pd.DataFrame(filas)
    tabla["mejor_tareas_reasignadas"] = False
    filas_escenario_c = tabla.index[tabla["escenario"] == "C - Falla"]
    if len(filas_escenario_c):
        minimo_reasignadas = tabla.loc[
            filas_escenario_c, "tareas_reasignadas_promedio"
        ].min()
        ganadoras = filas_escenario_c[
            tabla.loc[filas_escenario_c, "tareas_reasignadas_promedio"]
            == minimo_reasignadas
        ]
        tabla.loc[ganadoras, "mejor_tareas_reasignadas"] = True
    return tabla


def generar_csv(resultados: ResultadosPorCombinacion, ruta: str | Path) -> Path:
    """Escribe la tabla resumen de resultados en CSV."""
    destino = Path(ruta)
    destino.parent.mkdir(parents=True, exist_ok=True)
    tabla = _resumen(resultados)
    tabla.to_csv(destino, index=False, encoding="utf-8-sig")
    logger.info("CSV generado: %s", destino)
    return destino


def generar_excel(resultados: ResultadosPorCombinacion, ruta: str | Path) -> Path:
    """Escribe un libro Excel con una hoja por escenario."""
    destino = Path(ruta)
    destino.parent.mkdir(parents=True, exist_ok=True)
    tabla = _resumen(resultados)
    with pd.ExcelWriter(destino, engine="openpyxl") as writer:
        for escenario, datos in tabla.groupby("escenario", sort=False):
            hoja = escenario.split(" - ", maxsplit=1)[0]
            datos.to_excel(writer, sheet_name=hoja[:31], index=False)
            if escenario == "C - Falla":
                hoja_excel = writer.sheets[hoja[:31]]
                columna = datos.columns.get_loc("tareas_reasignadas_promedio") + 1
                resaltado = PatternFill(
                    fill_type="solid",
                    fgColor="D9EAD3",
                )
                for fila_excel, es_mejor in enumerate(
                    datos["mejor_tareas_reasignadas"], start=2
                ):
                    if es_mejor:
                        hoja_excel.cell(row=fila_excel, column=columna).fill = resaltado
    logger.info("Excel generado: %s", destino)
    return destino


def generar_graficas(
    resultados: ResultadosPorCombinacion,
    ruta: str | Path,
) -> list[Path]:
    """Genera gráficas de carga por servidor, makespan y balance de carga."""
    directorio = Path(ruta)
    directorio.mkdir(parents=True, exist_ok=True)
    tabla = _resumen(resultados)
    archivos: list[Path] = []

    for escenario, datos in tabla.groupby("escenario", sort=False):
        ids = sorted(
            int(columna.removeprefix("carga_servidor_").removesuffix("_promedio"))
            for columna in datos.columns
            if columna.startswith("carga_servidor_") and columna.endswith("_promedio")
        )
        fig, ax = plt.subplots(figsize=(9, 5))
        ancho = 0.8 / max(len(datos), 1)
        posiciones = list(range(len(ids)))
        for indice, (_, fila) in enumerate(datos.iterrows()):
            valores = [
                fila.get(f"carga_servidor_{servidor}_promedio", 0)
                for servidor in ids
            ]
            desplazamiento = (indice - (len(datos) - 1) / 2) * ancho
            ax.bar(
                [posicion + desplazamiento for posicion in posiciones],
                valores,
                width=ancho,
                label=fila["estrategia"],
            )
        ax.set(
            title=f"Carga por servidor — {escenario}",
            xlabel="Servidor",
            ylabel="Trabajo total asignado",
            xticks=posiciones,
            xticklabels=[f"S{id_}" for id_ in ids],
        )
        ax.legend(title="Estrategia", loc="upper left", bbox_to_anchor=(1.02, 1))
        fig.tight_layout(rect=(0, 0, 0.82, 1))
        destino = directorio / f"carga_{escenario.split(' - ')[0].lower()}.png"
        fig.savefig(destino, dpi=150)
        plt.close(fig)
        archivos.append(destino)

    for columna, titulo, eje_y, nombre in (
        (
            "makespan_promedio",
            "Makespan promedio por estrategia",
            "Unidades de tiempo",
            "makespan",
        ),
        (
            "desviacion_estandar_carga_promedio",
            "Desviación estándar de carga por estrategia",
            "Desviación estándar",
            "desviacion_carga",
        ),
        (
            "tiempo_espera_promedio_promedio",
            "Tiempo promedio de espera por estrategia",
            "Ticks de espera",
            "tiempo_espera",
        ),
        (
            "sobrecargas_total_promedio",
            "Sobrecargas totales por estrategia",
            "Eventos de sobrecarga",
            "sobrecargas",
        ),
    ):
        fig, ax = plt.subplots(figsize=(10, 6))
        escenarios = list(tabla["escenario"].drop_duplicates())
        estrategias = list(tabla["estrategia"].drop_duplicates())
        ancho = 0.8 / max(len(estrategias), 1)
        posiciones = list(range(len(escenarios)))
        for indice, estrategia in enumerate(estrategias):
            datos = tabla[tabla["estrategia"] == estrategia].set_index("escenario")
            valores = [datos.loc[escenario, columna] for escenario in escenarios]
            desplazamiento = (indice - (len(estrategias) - 1) / 2) * ancho
            ax.bar(
                [posicion + desplazamiento for posicion in posiciones],
                valores,
                width=ancho,
                label=estrategia,
            )
        ax.set(
            title=titulo,
            xlabel="Escenario",
            ylabel=eje_y,
            xticks=posiciones,
            xticklabels=escenarios,
        )
        ax.legend(title="Estrategia", loc="upper left", bbox_to_anchor=(1.02, 1))
        fig.tight_layout(rect=(0, 0, 0.82, 1))
        destino = directorio / f"{nombre}_comparativo.png"
        fig.savefig(destino, dpi=150)
        plt.close(fig)
        archivos.append(destino)

    logger.info("Gráficas generadas: %s", directorio)
    return archivos
