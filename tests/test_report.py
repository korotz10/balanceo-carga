import pandas as pd
from openpyxl import load_workbook

from src.dashboard import _estilos_mejores
from src.parameters import ParametrosSimulacion
from src.report import generar_csv, generar_excel, generar_graficas
from src.scenarios import ejecutar_todos


def test_dashboard_only_highlights_reassignment_metric_in_failure_scenario():
    tabla = pd.DataFrame(
        {
            "escenario": ["A - Normal", "A - Normal", "C - Falla", "C - Falla"],
            "tareas_reasignadas_promedio": [0.0, 0.0, 2.0, 1.0],
        }
    )

    estilos = _estilos_mejores(tabla)

    assert estilos.loc[0, "tareas_reasignadas_promedio"] == ""
    assert estilos.loc[1, "tareas_reasignadas_promedio"] == ""
    assert estilos.loc[2, "tareas_reasignadas_promedio"] == ""
    assert estilos.loc[3, "tareas_reasignadas_promedio"] == "background-color: #5353ec"


def test_reports_mark_reassignments_only_in_failure_scenario_and_add_comparisons(
    tmp_path,
):
    parametros = ParametrosSimulacion(
        num_tareas=12,
        num_servidores=4,
        duracion_min=1,
        duracion_max=4,
        intervalo_llegada=2,
        umbral_sobrecarga=12,
        tarea_falla=5,
        servidor_falla=2,
        servidor_lento=2,
        capacidad_reducida=0.5,
        semilla=24,
        repeticiones=2,
    )
    resultados = ejecutar_todos(parametros)
    csv_path = generar_csv(resultados, tmp_path / "resumen.csv")
    excel_path = generar_excel(resultados, tmp_path / "resumen.xlsx")
    graficas = generar_graficas(resultados, tmp_path / "graficas")

    tabla = pd.read_csv(csv_path)
    assert tabla.loc[tabla["escenario"] != "C - Falla", "mejor_tareas_reasignadas"].eq(
        False
    ).all()
    assert tabla.loc[tabla["escenario"] == "C - Falla", "mejor_tareas_reasignadas"].any()

    libro = load_workbook(excel_path)
    hoja_c = libro["C"]
    encabezados = [celda.value for celda in hoja_c[1]]
    columna_reasignadas = encabezados.index("tareas_reasignadas_promedio") + 1
    columna_mejor = encabezados.index("mejor_tareas_reasignadas") + 1
    for fila in range(2, hoja_c.max_row + 1):
        mejor = hoja_c.cell(row=fila, column=columna_mejor).value
        relleno = hoja_c.cell(row=fila, column=columna_reasignadas).fill
        assert (relleno.fill_type == "solid") is mejor

    nombres_graficas = {grafica.name for grafica in graficas}
    assert "tiempo_espera_comparativo.png" in nombres_graficas
    assert "sobrecargas_comparativo.png" in nombres_graficas
